"""Generate NOVA AI Technical Presentation PPTX using python-pptx."""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Color palette
BG = RGBColor(0x0F, 0x0B, 0x1E)
ACCENT = RGBColor(0x7C, 0x4D, 0xFF)
ACCENT2 = RGBColor(0xB3, 0x88, 0xFF)
ACCENT3 = RGBColor(0xE0, 0x40, 0xFB)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT = RGBColor(0xE0, 0xE0, 0xE0)
DIM = RGBColor(0xA0, 0xA0, 0xA0)
DARK_BG = RGBColor(0x1A, 0x1A, 0x2E)


def set_bg(slide, color=BG):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_text(slide, left, top, width, height, text, font_size=18, color=LIGHT,
             bold=False, alignment=PP_ALIGN.LEFT, font_name="Calibri"):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = alignment
    return txBox


def add_bullet_list(slide, left, top, width, height, items, font_size=16, color=LIGHT, spacing=Pt(6)):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = item
        p.font.size = Pt(font_size)
        p.font.color.rgb = color
        p.font.name = "Calibri"
        p.space_after = spacing
        p.level = 0
    return txBox


def add_code_block(slide, left, top, width, height, code_text, font_size=12):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top),
        Inches(width), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = DARK_BG
    shape.line.fill.background()
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.2)
    tf.margin_top = Inches(0.15)
    p = tf.paragraphs[0]
    p.text = code_text
    p.font.size = Pt(font_size)
    p.font.color.rgb = RGBColor(0xB3, 0x88, 0xFF)
    p.font.name = "Consolas"
    return shape


def add_table(slide, left, top, width, rows_data, col_widths, header_color=ACCENT):
    rows = len(rows_data)
    cols = len(col_widths)
    table_shape = slide.shapes.add_table(rows, cols, Inches(left), Inches(top),
                                          Inches(width), Inches(0.4 * rows))
    table = table_shape.table

    for i, w in enumerate(col_widths):
        table.columns[i].width = Inches(w)

    for row_idx, row_data in enumerate(rows_data):
        for col_idx, cell_text in enumerate(row_data):
            cell = table.cell(row_idx, col_idx)
            cell.text = cell_text
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.size = Pt(13)
                paragraph.font.name = "Calibri"
                if row_idx == 0:
                    paragraph.font.color.rgb = WHITE
                    paragraph.font.bold = True
                else:
                    paragraph.font.color.rgb = LIGHT
            if row_idx == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = header_color
            elif row_idx % 2 == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(0x1A, 0x1A, 0x3E)
            else:
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(0x14, 0x10, 0x28)
    return table_shape


# ============================================================
# SLIDE 1: Title
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
set_bg(slide)
add_text(slide, 1, 1.5, 11.3, 1.5, "NOVA AI", font_size=54, color=ACCENT,
         bold=True, alignment=PP_ALIGN.CENTER)
add_text(slide, 1, 3.2, 11.3, 1, "Local-First AI Agent Framework",
         font_size=28, color=ACCENT2, alignment=PP_ALIGN.CENTER)
add_text(slide, 1, 4.5, 11.3, 0.8, "Personal AI, On Personal Devices.",
         font_size=20, color=DIM, alignment=PP_ALIGN.CENTER)
add_text(slide, 1, 5.8, 11.3, 0.6,
         "GitHub: Hamza35779/NOVA-AI  |  Docs: hamza35779.github.io/NOVA-AI",
         font_size=14, color=DIM, alignment=PP_ALIGN.CENTER)

# ============================================================
# SLIDE 2: Problem
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)
add_text(slide, 0.8, 0.4, 11.7, 0.8, "The Problem: Cloud-Dependent \"Personal\" AI",
         font_size=32, color=ACCENT, bold=True)

add_table(slide, 0.8, 1.5, 11.7, [
    ["Problem", "Impact", "NOVA AI Solution"],
    ["Privacy", "All data leaves your device", "Local inference by default"],
    ["Cost", "$0.01-$0.10+ per query", "Local = $0.00, smart routing"],
    ["Latency", "200-2000ms network round-trips", "On-device inference, no network"],
    ["Availability", "No offline; outages = no AI", "Works completely offline"],
    ["Sovereignty", "No control over data/models", "Full user control, open source"],
], [2.5, 4.5, 4.7])

add_text(slide, 0.8, 5.5, 11.7, 0.8,
         "Local models already handle most tasks. What's missing is the software stack "
         "to make local-first AI practical.",
         font_size=16, color=DIM)

# ============================================================
# SLIDE 3: Architecture
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)
add_text(slide, 0.8, 0.4, 11.7, 0.8, "Architecture: Five Core Primitives",
         font_size=32, color=ACCENT, bold=True)

add_code_block(slide, 0.8, 1.4, 7, 3.5,
    "+----------------------------------------------------+\n"
    "|                   LEARNING                          |\n"
    "|        (Trace-driven routing & rewards)             |\n"
    "+----------+----------+----------+-------------------+\n"
    "|INTELLIGENCE|  ENGINE  | AGENTS  |      MEMORY        |\n"
    "| Model     | Inference| Pluggable|  Persistent        |\n"
    "| Catalog   | Backends | Agents   |  Searchable Storage |\n"
    "+----------+----------+----------+-------------------+\n"
    "         |              |              |\n"
    "         +--------------+--------------+\n"
    "                        |\n"
    "                   [EventBus]\n"
    "              (Trace Collection)", font_size=11)

add_bullet_list(slide, 8.2, 1.4, 4.5, 3.5, [
    "Intelligence - Model catalog with runtime discovery",
    "Engine - Inference backends (Ollama, vLLM, Cloud)",
    "Agents - 8+ pluggable agent types",
    "Memory - 5 searchable storage backends",
    "Learning - Trace-driven feedback loop",
    "EventBus - Thread-safe pub/sub communication",
], font_size=14)

add_text(slide, 0.8, 5.5, 11.7, 0.8,
         "Feedback loop: agents produce traces -> traces inform learning -> "
         "learning improves routing -> better routing improves agent performance.",
         font_size=14, color=DIM)

# ============================================================
# SLIDE 4: Engine Layer
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)
add_text(slide, 0.8, 0.4, 11.7, 0.8, "Engine Layer: Pluggable Inference Runtime",
         font_size=32, color=ACCENT, bold=True)

add_code_block(slide, 0.8, 1.4, 5.5, 1.8,
    "class InferenceEngine(ABC):\n"
    "    def generate(messages, model) -> dict\n"
    "    def stream(messages, model) -> Iterator[str]\n"
    "    def list_models() -> list[str]\n"
    "    def health() -> bool", font_size=13)

add_table(slide, 0.8, 3.5, 11.7, [
    ["Backend", "Type", "Technology", "Use Case"],
    ["Ollama", "Local", "Native HTTP API", "Default local engine, broad model support"],
    ["vLLM", "Local", "OpenAI-compat API", "High-throughput, PagedAttention"],
    ["SGLang", "Local", "OpenAI-compat API", "Structured generation"],
    ["llama.cpp", "Local", "OpenAI-compat API", "GGUF quantized, CPU/GPU hybrid"],
    ["Cloud", "Cloud", "Provider SDKs", "OpenAI, Anthropic, Google, DeepSeek"],
], [2, 1.5, 3, 5.2])

add_text(slide, 0.8, 6.2, 11.7, 0.8,
         "Smart Router: model=\"auto\" scores complexity -> simple = local, complex = cloud. "
         "Cost tracking: local = $0.00, cloud = per-query USD.",
         font_size=14, color=DIM)

# ============================================================
# SLIDE 5: Agent Framework
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)
add_text(slide, 0.8, 0.4, 11.7, 0.8, "Agent Framework: 8+ Agent Types",
         font_size=32, color=ACCENT, bold=True)

add_code_block(slide, 0.8, 1.3, 5.5, 3.2,
    "BaseAgent (ABC)\n"
    "  +-- SimpleAgent          (single-turn)\n"
    "  +-- OpenHandsAgent       (openhands-sdk)\n"
    "  +-- ClaudeCodeAgent      (Claude SDK)\n"
    "  +-- SandboxedAgent       (Docker isolation)\n"
    "  +-- ToolUsingAgent       (accepts tools)\n"
    "        +-- OrchestratorAgent  (tool loop)\n"
    "        +-- NativeReActAgent   (ReAct)\n"
    "        +-- NativeOpenHandsAgent (CodeAct)\n"
    "        +-- RLMAgent           (recursive LM)\n"
    "        +-- OperativeAgent     (persistent)\n"
    "        +-- MonitorOperative   (long-horizon)", font_size=12)

add_table(slide, 6.8, 1.3, 5.7, [
    ["Agent", "Mode", "Description"],
    ["simple", "On-demand", "Single-turn Q&A, no tools"],
    ["orchestrator", "On-demand", "Multi-turn tool-calling loop (default)"],
    ["native_react", "On-demand", "Thought-Action-Observation loop"],
    ["native_openhands", "On-demand", "CodeAct - writes and executes Python"],
    ["deep_research", "On-demand", "Multi-hop research with citations"],
    ["operative", "Continuous", "Persistent agent with state"],
    ["monitor_operative", "Continuous", "Long-horizon monitoring"],
    ["proactive_agent", "Scheduled", "Autonomous routine tasks"],
], [2.5, 1.8, 2.4])

add_text(slide, 0.8, 5.8, 11.7, 0.8,
         "OrchestratorAgent (default): function_calling mode (OpenAI format) or structured mode "
         "(THOUGHT/TOOL/INPUT/FINAL_ANSWER) used by SFT/GRPO training pipelines.",
         font_size=14, color=DIM)

# ============================================================
# SLIDE 6: Tool System
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)
add_text(slide, 0.8, 0.4, 11.7, 0.8, "Tool System: 58+ Built-in Tools",
         font_size=32, color=ACCENT, bold=True)

add_code_block(slide, 0.8, 1.3, 5.5, 1.5,
    "class BaseTool(ABC):\n"
    "    tool_id: str\n"
    "    def spec(self) -> ToolSpec\n"
    "    def execute(**params) -> ToolResult\n"
    "    def to_openai_function() -> Dict", font_size=12)

add_table(slide, 0.8, 3.0, 11.7, [
    ["Category", "Tools"],
    ["Web", "web_search, browser_navigate, browser_click, browser_extract, http_request"],
    ["Code", "code_interpreter, code_interpreter_docker, shell_exec, repl, code_scaffolder"],
    ["Files", "file_read, file_write, apply_patch, git_tool, git_manager, file_converter"],
    ["Knowledge", "knowledge_search, retrieval, scan_chunks, memory_manage, memory_wiki_tools"],
    ["Media", "image_tool, audio_tool, text_to_speech, screen_capture, screen_monitor"],
    ["System", "system_monitor, scheduler_tool, canvas_tool, clipboard_ai, api_tester"],
    ["Integration", "mcp_adapter, docker_shell_exec, cisco_packet_tracer, channel_tools"],
], [2, 9.7])

# ============================================================
# SLIDE 7: Memory
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)
add_text(slide, 0.8, 0.4, 11.7, 0.8, "Memory: 5 Searchable Storage Backends",
         font_size=32, color=ACCENT, bold=True)

add_table(slide, 0.8, 1.4, 11.7, [
    ["Backend", "Technology", "Best For", "Dependencies"],
    ["SQLite/FTS5", "Full-text search", "Quick start, zero config", "None (built-in)"],
    ["FAISS", "Dense vectors", "Semantic search", "sentence-transformers, faiss-cpu"],
    ["ColBERTv2", "Late interaction", "High-precision retrieval", "colbert-ai, torch"],
    ["BM25", "Term-frequency", "Keyword search", "rank-bm25"],
    ["Hybrid", "Reciprocal Rank Fusion", "Best of sparse + dense", "FAISS + BM25"],
], [2.5, 3.5, 3, 4.7])

add_text(slide, 0.8, 4.5, 11.7, 2.5,
         "Ingestion Pipeline:\n"
         "  File reading -> Chunking (ChunkConfig) -> Embedding (SentenceTransformerEmbedder) -> Storage\n\n"
         "Context Injection:\n"
         "  When agent.context_from_memory is enabled, relevant documents are retrieved and prepended "
         "to the prompt with source attribution (file path, chunk index).",
         font_size=14, color=LIGHT)

# ============================================================
# SLIDE 8: Learning
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)
add_text(slide, 0.8, 0.4, 11.7, 0.8, "Learning: Trace-Driven Feedback",
         font_size=32, color=ACCENT, bold=True)

add_code_block(slide, 0.8, 1.3, 7, 2.5,
    "Agent runs query\n"
    "    -> Produces Trace (steps, tools, timing, cost)\n"
    "        -> TraceAnalyzer computes statistics\n"
    "            -> TraceDrivenPolicy learns optimal combinations\n"
    "                -> Better routing for future queries", font_size=14)

add_table(slide, 8.5, 1.3, 4.2, [
    ["Component", "Purpose"],
    ["TraceStore", "SQLite trace persistence"],
    ["TraceAnalyzer", "Aggregated statistics"],
    ["HeuristicRouter", "Complexity scoring"],
    ["TraceDrivenPolicy", "Learned routing"],
    ["SkillForge", "Skill synthesis"],
], [2, 2.2])

add_text(slide, 0.8, 4.2, 11.7, 3,
         "Training Pipelines:\n\n"
         "  SFT (Supervised Fine-Tuning) - trains on trace data\n"
         "  GRPO (Group Relative Policy Optimization) - RL-based training\n"
         "  DPO (Direct Preference Optimization) - preference-based training\n"
         "  LoRA (Low-Rank Adaptation) - efficient fine-tuning\n\n"
         "SkillForge Pipeline:\n\n"
         "  Miner -> Synthesizer -> Pipeline -> Gauntlet -> Store -> Adoption",
         font_size=14, color=LIGHT)

# ============================================================
# SLIDE 9: Security
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)
add_text(slide, 0.8, 0.4, 11.7, 0.8, "Security Model",
         font_size=32, color=ACCENT, bold=True)

add_table(slide, 0.8, 1.3, 11.7, [
    ["Component", "Mechanism", "Description"],
    ["Loop Guard", "SHA-256 hash tracking", "Detects degenerate tool-calling loops, ping-pong, per-tool budgets"],
    ["RBAC", "Capability policy", "Tools declare required_capabilities; agents need matching permissions"],
    ["Secret Scanner", "Pattern matching", "Detects hardcoded API keys and secrets in code"],
    ["PII Scanner", "Pattern matching", "Detects personally identifiable information"],
    ["Boundary Guard", "Argument scanning", "Scans external tool arguments for security threats"],
    ["SSRF Protection", "URL validation", "Web tools check for server-side request forgery"],
    ["Taint Checking", "Provenance tracking", "Tracks data provenance through tool pipelines"],
    ["Redaction", "PII/secret scrubbing", "Scrubs secrets/PII before cloud transmission"],
    ["Budget", "Cost/token limits", "Per-agent limits with automatic pause on exceed"],
    ["Circuit Breaker", "Failure counting", "Auto-pauses operator after consecutive failures"],
], [2.5, 3, 6.2])

# ============================================================
# SLIDE 10: Deployment
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)
add_text(slide, 0.8, 0.4, 11.7, 0.8, "Deployment Options",
         font_size=32, color=ACCENT, bold=True)

add_table(slide, 0.8, 1.3, 11.7, [
    ["Platform", "Method", "Requirements"],
    ["Windows", "Inno Setup installer (.exe)", "None - self-contained, no uv/git/Python"],
    ["Linux/macOS", "Install script (curl | bash)", "Python 3.10-3.13, uv"],
    ["Docker", "docker-compose.yml", "Docker + optional NVIDIA GPU"],
    ["Desktop", "Tauri app (.exe/.dmg/.AppImage)", "uv (auto-installed on first launch)"],
    ["PyPI", "pip install nova-ai-pro", "Python 3.10-3.13"],
    ["From Source", "git clone + uv sync", "Python 3.10-3.13, uv, git"],
], [2.5, 4, 5.2])

add_text(slide, 0.8, 5.2, 11.7, 2,
         "Channel Integrations (15+):\n\n"
         "Telegram, Discord, Slack, WhatsApp, Line, Viber, Messenger, Reddit, Mastodon, "
         "XMPP, Rocket.Chat, Zulip, Twitter/X, Twitch, Nostr, Twilio, Gmail",
         font_size=14, color=LIGHT)

# ============================================================
# SLIDE 11: Benchmarks
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)
add_text(slide, 0.8, 0.4, 11.7, 0.8, "Performance & Cost Tracking",
         font_size=32, color=ACCENT, bold=True)

add_table(slide, 0.8, 1.4, 11.7, [
    ["Metric", "Local Inference", "Cloud Inference"],
    ["Cost per query", "$0.00", "$0.01 - $0.10+"],
    ["Latency", "50-500ms (on-device)", "200-2000ms (network)"],
    ["Availability", "Always available", "Requires internet"],
    ["Privacy", "Data stays on device", "Data sent to provider"],
    ["Throughput", "Hardware dependent", "Provider rate limits"],
    ["Energy", "Tracked via pynvml/amdsmi", "N/A"],
], [3, 4.35, 4.35])

add_text(slide, 0.8, 5.2, 11.7, 2,
         "Benchmark Framework:\n\n"
         "  LatencyBenchmark: per-call latency measurement\n"
         "  ThroughputBenchmark: tokens/second measurement\n"
         "  Energy Monitoring: GPU power consumption via pynvml (NVIDIA), amdsmi (AMD), zeus-ml (Apple)\n"
         "  Cost Tracking: per-query USD via OpenRouter or pricing table estimate",
         font_size=14, color=LIGHT)

# ============================================================
# SLIDE 12: Roadmap
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)
add_text(slide, 0.8, 0.4, 11.7, 0.8, "Roadmap",
         font_size=32, color=ACCENT, bold=True)

add_text(slide, 0.8, 1.3, 5.5, 0.6, "Current (v1.2.4)", font_size=20, color=ACCENT2, bold=True)
add_bullet_list(slide, 0.8, 2.0, 5.5, 4, [
    "Five-primitive architecture",
    "8+ agent types",
    "58+ built-in tools",
    "5 memory backends",
    "Smart model router",
    "MCP integration",
    "Desktop app (Tauri)",
    "Windows/Linux/macOS installers",
    "15+ channel integrations",
    "Trace-driven learning",
    "Skill synthesis (SkillForge)",
    "SFT/GRPO/DPO training",
], font_size=14, color=LIGHT)

add_text(slide, 7, 1.3, 5.5, 0.6, "Planned", font_size=20, color=ACCENT3, bold=True)
add_bullet_list(slide, 7, 2.0, 5.5, 4, [
    "Enhanced hybrid local+cloud paradigms",
    "Improved skill validation & adoption",
    "Additional channel integrations",
    "Mobile companion app",
    "Enterprise deployment features",
    "Advanced energy monitoring",
    "Orchestrator-8B trained variant",
    "SkillOrchestra learn-phase",
    "Advisor RL training",
], font_size=14, color=LIGHT)

# ============================================================
# SLIDE 13: Call to Action
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)
add_text(slide, 1, 1.5, 11.3, 1.5, "Get Started", font_size=48, color=ACCENT,
         bold=True, alignment=PP_ALIGN.CENTER)
add_text(slide, 1, 3.2, 11.3, 1, "Personal AI, On Personal Devices.",
         font_size=24, color=ACCENT2, alignment=PP_ALIGN.CENTER)

add_code_block(slide, 2.5, 4.2, 8.3, 1.5,
    "# Install (Linux/macOS)\n"
    "curl -fsSL https://hamza35779.github.io/NOVA-AI/install.sh | bash\n\n"
    "# Install (Windows)\n"
    "irm https://hamza35779.github.io/NOVA-AI/install.ps1 | iex", font_size=14)

add_text(slide, 1, 6.2, 11.3, 0.6,
         "GitHub: Hamza35779/NOVA-AI  |  Docs: hamza35779.github.io/NOVA-AI  |  License: Apache 2.0",
         font_size=14, color=DIM, alignment=PP_ALIGN.CENTER)

# === SAVE ===
output_path = r"D:\My Softwares\Friday AI\NOVA AI\docs\proposal\NOVA_AI_Technical_Presentation.pptx"
prs.save(output_path)
print(f"PPTX saved to: {output_path}")
